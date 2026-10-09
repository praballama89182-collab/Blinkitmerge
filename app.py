import os
import io
import pandas as pd
import numpy as np
import streamlit as st
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

st.set_page_config(
    page_title="Blinkit Report Merger & Analytics",
    page_icon="📊",
    layout="wide"
)

# Professional blue dashboard theme + clean, centered data presentation
st.markdown("""
<style>
    :root {
        --dash-blue: #4A90C2;
        --dash-blue-dark: #285B7A;
        --dash-blue-light: #EEF6FC;
        --dash-blue-soft: #DCECF8;
        --dash-border: #C8DCEB;
        --dash-text: #29465B;
    }

    .stApp {
        background: linear-gradient(180deg, #F6FAFD 0%, #FFFFFF 38%);
        color: var(--dash-text);
    }

    [data-testid="stHeader"] {
        background: rgba(255,255,255,0.92);
    }

    h1, h2, h3 {
        color: var(--dash-blue-dark) !important;
        font-weight: 700 !important;
    }

    [data-testid="stMetric"] {
        background: linear-gradient(135deg, #FFFFFF 0%, #EEF6FF 100%);
        border: 1px solid var(--dash-border);
        border-left: 5px solid #6AA6CF;
        border-radius: 12px;
        padding: 14px 18px;
        box-shadow: 0 3px 12px rgba(74, 144, 194, 0.10);
    }

    [data-testid="stMetricLabel"] {
        color: #52708D !important;
        font-weight: 600 !important;
    }

    [data-testid="stMetricValue"] {
        color: var(--dash-blue-dark) !important;
        font-weight: 750 !important;
    }

    .stTabs [data-baseweb="tab-list"] {
        gap: 5px;
        background: #F0F7FC;
        border: 1px solid var(--dash-border);
        border-radius: 12px;
        padding: 5px;
    }

    .stTabs [data-baseweb="tab"] {
        height: 42px;
        border-radius: 7px;
        color: #315A7D;
        font-weight: 600;
    }

    .stTabs [aria-selected="true"] {
        background: linear-gradient(135deg, #6AA6CF 0%, #4A90C2 100%) !important;
        color: #FFFFFF !important;
        box-shadow: 0 2px 7px rgba(74, 144, 194, 0.20);
    }

    [data-testid="stFileUploader"] {
        background: #FFFFFF;
        border: 1px solid var(--dash-border);
        border-radius: 12px;
        padding: 8px;
        box-shadow: 0 2px 8px rgba(74, 144, 194, 0.07);
    }

    [data-testid="stFileUploaderDropzone"] {
        background: #F7FBFE;
        border: 1px dashed #9FC4DE;
        border-radius: 9px;
    }

    [data-testid="stFileUploader"] button,
    .stDownloadButton button,
    .stButton button {
        background: linear-gradient(135deg, #6AA6CF, #4A90C2) !important;
        color: #FFFFFF !important;
        border: 0 !important;
        border-radius: 8px !important;
        font-weight: 650 !important;
        box-shadow: 0 2px 7px rgba(11, 92, 173, 0.20);
    }

    [data-testid="stFileUploader"] button:hover,
    .stDownloadButton button:hover,
    .stButton button:hover {
        background: #397BAA !important;
        color: #FFFFFF !important;
    }

    [data-testid="stDataFrame"] {
        border: 1px solid var(--dash-border);
        border-radius: 10px;
        overflow: hidden;
        box-shadow: 0 2px 9px rgba(11, 92, 173, 0.06);
    }

    [data-testid="stDataFrame"] th,
    [data-testid="stDataFrame"] td {
        text-align: center !important;
        vertical-align: middle !important;
    }

    [data-testid="stDataFrame"] th {
        background: #E5F1F9 !important;
        color: #285B7A !important;
        font-weight: 700 !important;
    }

    [data-baseweb="select"] > div {
        border-color: #AFC7DF !important;
        border-radius: 8px !important;
    }

    [data-testid="stRadio"] label,
    [data-testid="stSelectbox"] label,
    [data-testid="stMultiSelect"] label {
        color: #315A7D !important;
        font-weight: 600 !important;
    }

    hr {
        border-color: #D7E4F2 !important;
    }
</style>
""", unsafe_allow_html=True)

st.title("📊 Blinkit Ad Report Merger & Analytics")
st.write("Upload up to 5 monthly Excel ad campaign spreadsheets (.xlsx, .xls, .xlsb, .xlsm). Preview raw & consolidated sheets, view Month-on-Month Comparison tables, and analyze multi-metric curve & bar trend comparisons.")

def style_and_export_pivot(pivot_df, sheet_name="Comparison"):
    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    
    top_header_fill = PatternFill(start_color="4A90C2", end_color="4A90C2", fill_type="solid")
    top_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    sec_header_fill = PatternFill(start_color="E5F1F9", end_color="E5F1F9", fill_type="solid")
    sec_header_font = Font(name="Calibri", size=10, bold=True, color="285B7A")
    
    index_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    index_font = Font(name="Calibri", size=10, bold=True, color="000000")
    
    total_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    total_font = Font(name="Calibri", size=10, bold=True, color="000000")
    
    data_font = Font(name="Calibri", size=10, color="000000")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    if isinstance(pivot_df.columns, pd.MultiIndex):
        metrics = pivot_df.columns.get_level_values(1)
        entity_title = pivot_df.index.name if pivot_df.index.name else ""

        ws.cell(row=1, column=1, value="").fill = top_header_fill
        ws.cell(row=1, column=1).border = thin_border
        
        current_col = 2
        for month in pivot_df.columns.get_level_values(0).unique():
            sub_cols = [c for c in pivot_df.columns if c[0] == month]
            num_sub = len(sub_cols)
            if num_sub > 1:
                ws.merge_cells(start_row=1, start_column=current_col, end_row=1, end_column=current_col + num_sub - 1)
            cell = ws.cell(row=1, column=current_col, value=str(month))
            cell.fill = top_header_fill
            cell.font = top_header_font
            cell.alignment = align_center
            
            for c in range(current_col, current_col + num_sub):
                ws.cell(row=1, column=c).border = thin_border
                ws.cell(row=1, column=c).fill = top_header_fill
            current_col += num_sub

        cell_a2 = ws.cell(row=2, column=1, value=str(entity_title))
        cell_a2.fill = index_fill
        cell_a2.font = index_font
        cell_a2.alignment = align_center
        cell_a2.border = thin_border

        for col_idx, metric in enumerate(metrics, start=2):
            cell = ws.cell(row=2, column=col_idx, value=str(metric))
            cell.fill = sec_header_fill
            cell.font = sec_header_font
            cell.alignment = align_center
            cell.border = thin_border

        start_data_row = 3
        for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=start_data_row):
            is_grand_total = (str(idx_val).strip().lower() == 'grand total')
            
            idx_cell = ws.cell(row=r_idx, column=1, value=str(idx_val))
            idx_cell.fill = total_fill if is_grand_total else index_fill
            idx_cell.font = total_font if is_grand_total else index_font
            idx_cell.alignment = align_left
            idx_cell.border = thin_border

            for c_idx, (m_col, val) in enumerate(zip(metrics, row_data), start=2):
                val_cell = ws.cell(row=r_idx, column=c_idx)
                val_cell.font = total_font if is_grand_total else data_font
                if is_grand_total:
                    val_cell.fill = total_fill
                val_cell.alignment = align_center
                val_cell.border = thin_border

                if isinstance(val, (int, float, np.number)):
                    if str(m_col).upper() == 'ACOS':
                        val_cell.value = val / 100.0 if val > 1 else val
                        val_cell.number_format = '0.00%'
                    elif str(m_col).upper() == 'CPM':
                        val_cell.value = round(val)
                        val_cell.number_format = '#,##0'
                    else:
                        val_cell.value = round(val, 2)
                        val_cell.number_format = '#,##0' if float(val).is_integer() else '#,##0.00'
                else:
                    val_cell.value = val
    else:
        for col_idx, col_name in enumerate(pivot_df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=str(col_name))
            cell.fill = top_header_fill
            cell.font = top_header_font
            cell.alignment = align_center
            cell.border = thin_border

        for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=2):
            is_pct_row = (str(row_data.iloc[0]).strip().lower() == 'percentage %')
            
            for c_idx, (col_name, val) in enumerate(zip(pivot_df.columns, row_data), start=1):
                val_cell = ws.cell(row=r_idx, column=c_idx)
                val_cell.font = total_font if is_pct_row else data_font
                val_cell.alignment = align_center
                val_cell.border = thin_border
                
                if str(col_name).upper() == 'ACOS' and not is_pct_row and isinstance(val, (int, float, np.number)):
                    val_cell.value = val / 100.0 if val > 1 else val
                    val_cell.number_format = '0.00%'
                elif str(col_name).upper() == 'CPM' and not is_pct_row and isinstance(val, (int, float, np.number)):
                    val_cell.value = round(val)
                    val_cell.number_format = '#,##0'
                elif isinstance(val, float):
                    val_cell.value = round(val, 2)
                    val_cell.number_format = '#,##0' if float(val).is_integer() else '#,##0.00'
                elif isinstance(val, (int, np.integer)):
                    val_cell.value = val
                    val_cell.number_format = '#,##0'
                else:
                    val_cell.value = val

                if is_pct_row and c_idx > 1:
                    val_str = str(val).replace('%', '').replace('+', '').strip()
                    try:
                        num_v = float(val_str)
                        if num_v < 0:
                            val_cell.fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
                            val_cell.font = Font(name="Calibri", size=10, bold=True, color="721C24")
                        elif num_v > 0:
                            val_cell.fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
                            val_cell.font = Font(name="Calibri", size=10, bold=True, color="155724")
                    except ValueError:
                        pass
                elif is_pct_row:
                    val_cell.fill = sec_header_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def convert_all_pivots_to_excel(pivot_dict):
    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    top_header_fill = PatternFill(start_color="4A90C2", end_color="4A90C2", fill_type="solid")
    top_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sec_header_fill = PatternFill(start_color="E5F1F9", end_color="E5F1F9", fill_type="solid")
    sec_header_font = Font(name="Calibri", size=10, bold=True, color="285B7A")
    index_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    index_font = Font(name="Calibri", size=10, bold=True, color="000000")
    total_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    total_font = Font(name="Calibri", size=10, bold=True, color="000000")
    data_font = Font(name="Calibri", size=10, color="000000")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    for sheet_name, pivot_df in pivot_dict.items():
        if pivot_df is not None and not pivot_df.empty:
            ws = wb.create_sheet(title=sheet_name[:31])

            if isinstance(pivot_df.columns, pd.MultiIndex):
                metrics = pivot_df.columns.get_level_values(1)
                entity_title = pivot_df.index.name if pivot_df.index.name else ""

                ws.cell(row=1, column=1, value="").fill = top_header_fill
                ws.cell(row=1, column=1).border = thin_border
                
                current_col = 2
                for month in pivot_df.columns.get_level_values(0).unique():
                    sub_cols = [c for c in pivot_df.columns if c[0] == month]
                    num_sub = len(sub_cols)
                    if num_sub > 1:
                        ws.merge_cells(start_row=1, start_column=current_col, end_row=1, end_column=current_col + num_sub - 1)
                    cell = ws.cell(row=1, column=current_col, value=str(month))
                    cell.fill = top_header_fill
                    cell.font = top_header_font
                    cell.alignment = align_center
                    
                    for c in range(current_col, current_col + num_sub):
                        ws.cell(row=1, column=c).border = thin_border
                        ws.cell(row=1, column=c).fill = top_header_fill
                    current_col += num_sub

                cell_a2 = ws.cell(row=2, column=1, value=str(entity_title))
                cell_a2.fill = index_fill
                cell_a2.font = index_font
                cell_a2.alignment = align_center
                cell_a2.border = thin_border

                for col_idx, metric in enumerate(metrics, start=2):
                    cell = ws.cell(row=2, column=col_idx, value=str(metric))
                    cell.fill = sec_header_fill
                    cell.font = sec_header_font
                    cell.alignment = align_center
                    cell.border = thin_border

                for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=3):
                    is_grand_total = (str(idx_val).strip().lower() == 'grand total')
                    
                    idx_cell = ws.cell(row=r_idx, column=1, value=str(idx_val))
                    idx_cell.fill = total_fill if is_grand_total else index_fill
                    idx_cell.font = total_font if is_grand_total else index_font
                    idx_cell.alignment = align_left
                    idx_cell.border = thin_border

                    for c_idx, (m_col, val) in enumerate(zip(metrics, row_data), start=2):
                        val_cell = ws.cell(row=r_idx, column=c_idx)
                        val_cell.font = total_font if is_grand_total else data_font
                        if is_grand_total:
                            val_cell.fill = total_fill
                        val_cell.alignment = align_center
                        val_cell.border = thin_border

                        if isinstance(val, (int, float, np.number)):
                            if str(m_col).upper() == 'ACOS':
                                val_cell.value = val / 100.0 if val > 1 else val
                                val_cell.number_format = '0.00%'
                            elif str(m_col).upper() == 'CPM':
                                val_cell.value = round(val)
                                val_cell.number_format = '#,##0'
                            else:
                                val_cell.value = round(val, 2)
                                val_cell.number_format = '#,##0' if float(val).is_integer() else '#,##0.00'
                        else:
                            val_cell.value = val
            else:
                for col_idx, col_name in enumerate(pivot_df.columns, start=1):
                    cell = ws.cell(row=1, column=col_idx, value=str(col_name))
                    cell.fill = top_header_fill
                    cell.font = top_header_font
                    cell.alignment = align_center
                    cell.border = thin_border

                for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=2):
                    is_pct_row = (str(row_data.iloc[0]).strip().lower() == 'percentage %')
                    for c_idx, (col_name, val) in enumerate(zip(pivot_df.columns, row_data), start=1):
                        val_cell = ws.cell(row=r_idx, column=c_idx)
                        val_cell.font = total_font if is_pct_row else data_font
                        val_cell.alignment = align_center
                        val_cell.border = thin_border
                        
                        if str(col_name).upper() == 'ACOS' and not is_pct_row and isinstance(val, (int, float, np.number)):
                            val_cell.value = val / 100.0 if val > 1 else val
                            val_cell.number_format = '0.00%'
                        elif str(col_name).upper() == 'CPM' and not is_pct_row and isinstance(val, (int, float, np.number)):
                            val_cell.value = round(val)
                            val_cell.number_format = '#,##0'
                        elif isinstance(val, float):
                            val_cell.value = round(val, 2)
                            val_cell.number_format = '#,##0.00'
                        elif isinstance(val, (int, np.integer)):
                            val_cell.value = val
                            val_cell.number_format = '#,##0'
                        else:
                            val_cell.value = val

                        if is_pct_row and c_idx > 1:
                            val_str = str(val).replace('%', '').replace('+', '').strip()
                            try:
                                num_v = float(val_str)
                                if num_v < 0:
                                    val_cell.fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
                                    val_cell.font = Font(name="Calibri", size=10, bold=True, color="721C24")
                                elif num_v > 0:
                                    val_cell.fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
                                    val_cell.font = Font(name="Calibri", size=10, bold=True, color="155724")
                            except ValueError:
                                pass
                        elif is_pct_row:
                            val_cell.fill = sec_header_fill

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# File Upload Section
col1, col2, col3, col4, col5 = st.columns(5)
uploaded_files = []

with col1:
    f1 = st.file_uploader("Upload File 1", type=["xlsx", "xls", "xlsb", "xlsm"], key="file1")
    if f1: uploaded_files.append(f1)

with col2:
    f2 = st.file_uploader("Upload File 2", type=["xlsx", "xls", "xlsb", "xlsm"], key="file2")
    if f2: uploaded_files.append(f2)

with col3:
    f3 = st.file_uploader("Upload File 3", type=["xlsx", "xls", "xlsb", "xlsm"], key="file3")
    if f3: uploaded_files.append(f3)

with col4:
    f4 = st.file_uploader("Upload File 4", type=["xlsx", "xls", "xlsb", "xlsm"], key="file4")
    if f4: uploaded_files.append(f4)

with col5:
    f5 = st.file_uploader("Upload File 5", type=["xlsx", "xls", "xlsb", "xlsm"], key="file5")
    if f5: uploaded_files.append(f5)

if uploaded_files:
    consolidated_dfs = []
    raw_files_dict = {}
    consolidated_raw_tabs = {}

    for uploaded_file in uploaded_files:
        fallback_month_name = os.path.splitext(uploaded_file.name)[0].upper()
        raw_files_dict[uploaded_file.name] = {}
        
        try:
            xls = pd.ExcelFile(uploaded_file)
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name)
                raw_files_dict[uploaded_file.name][sheet_name] = df.copy()
                
                date_col = None
                for col_candidate in ['Date', 'date', 'Day', 'DATE']:
                    if col_candidate in df.columns:
                        date_col = col_candidate
                        break

                if date_col:
                    dt_series = pd.to_datetime(df[date_col], dayfirst=True, errors='coerce')
                    month_series = dt_series.dt.strftime('%B').str.upper()
                    df['Month'] = month_series.fillna(fallback_month_name)
                else:
                    df['Month'] = fallback_month_name

                df_raw = df.copy()
                if sheet_name not in consolidated_raw_tabs:
                    consolidated_raw_tabs[sheet_name] = []
                consolidated_raw_tabs[sheet_name].append(df_raw)
                
                df_consolidated = df.copy()
                if 'Month' in df_consolidated.columns:
                    col_month = df_consolidated.pop('Month')
                    df_consolidated.insert(0, 'Month', col_month)
                else:
                    df_consolidated.insert(0, 'Month', fallback_month_name)
                    
                df_consolidated.insert(1, 'Tab Name', sheet_name)
                consolidated_dfs.append(df_consolidated)
        except Exception as e:
            st.error(f"Error processing {uploaded_file.name}: {str(e)}")

    if consolidated_dfs:
        final_df = pd.concat(consolidated_dfs, ignore_index=True)
        
        base_cols = ['Month', 'Tab Name']
        product_listing_cols = []
        for df in consolidated_dfs:
            if 'PRODUCT_LISTING' in df['Tab Name'].values:
                pl_df = df[df['Tab Name'] == 'PRODUCT_LISTING']
                product_listing_cols = [c for c in pl_df.columns if c not in base_cols]
                break
        
        remaining_cols = [c for c in final_df.columns if c not in base_cols and c not in product_listing_cols]
        final_df = final_df.reindex(columns=base_cols + product_listing_cols + remaining_cols)

        match_col = 'Match Type' if 'Match Type' in final_df.columns else None
        target_col = 'Targeting Type' if 'Targeting Type' in final_df.columns else None
        tab_col = 'Tab Name' if 'Tab Name' in final_df.columns else None

        def is_na_series(series):
            if series is None or series.empty:
                return pd.Series(True, index=final_df.index)
            cleaned = series.astype(str).str.strip().str.upper()
            return series.isna() | cleaned.isin(['NA', 'N/A', 'NAN', 'NONE', ''])

        final_df['_filled_fallback'] = False

        if match_col:
            na_match = is_na_series(final_df[match_col])
            if target_col:
                valid_target = ~is_na_series(final_df[target_col])
                fill_from_target = na_match & valid_target
                final_df.loc[fill_from_target, match_col] = final_df.loc[fill_from_target, target_col]
                final_df.loc[fill_from_target, '_filled_fallback'] = True
                na_match = is_na_series(final_df[match_col])

            if tab_col:
                fill_from_tab = na_match
                final_df.loc[fill_from_tab, match_col] = final_df.loc[fill_from_tab, tab_col]
                final_df.loc[fill_from_tab, '_filled_fallback'] = True
        elif target_col:
            final_df['Match Type'] = final_df[target_col]
            na_match = is_na_series(final_df['Match Type'])
            if tab_col:
                final_df.loc[na_match, 'Match Type'] = final_df.loc[na_match, tab_col]
                final_df.loc[na_match, '_filled_fallback'] = True
            match_col = 'Match Type'

        def get_numeric_col(df, possible_cols):
            for col in possible_cols:
                if col in df.columns:
                    return pd.to_numeric(df[col], errors='coerce').fillna(0)
            return pd.Series(0, index=df.index)

        final_df['_impressions'] = get_numeric_col(final_df, ['Impressions'])
        final_df['_direct_atc'] = get_numeric_col(final_df, ['Direct ATC'])
        final_df['_indirect_atc'] = get_numeric_col(final_df, ['Indirect ATC'])
        final_df['_atc'] = final_df['_direct_atc'] + final_df['_indirect_atc']

        final_df['_direct_orders'] = get_numeric_col(final_df, ['Direct Quantities Sold', 'Direct Orders'])
        final_df['_indirect_orders'] = get_numeric_col(final_df, ['Indirect Quantities Sold', 'Indirect Orders'])
        final_df['_orders'] = final_df['_direct_orders'] + final_df['_indirect_orders']

        final_df['_direct_sales'] = get_numeric_col(final_df, ['Direct Sales'])
        final_df['_indirect_sales'] = get_numeric_col(final_df, ['Indirect Sales'])
        final_df['_sales'] = final_df['_direct_sales'] + final_df['_indirect_sales']

        final_df['_budget_consumed'] = get_numeric_col(final_df, ['Estimated Budget Consumed', 'Budget Consumed', 'Spend'])

        if match_col and match_col in final_df.columns:
            final_df['Ad Type Combined'] = final_df[match_col].fillna("Other").astype(str).str.title()
        else:
            final_df['Ad Type Combined'] = "Other"

        date_col = None
        for col_candidate in ['Date', 'date', 'Day', 'DATE']:
            if col_candidate in final_df.columns:
                date_col = col_candidate
                break

        if date_col:
            final_df['_date_dt'] = pd.to_datetime(final_df[date_col], dayfirst=True, errors='coerce')
            def assign_week_formatted(row):
                dt = row['_date_dt']
                if pd.isna(dt): return np.nan
                day = dt.day
                if 1 <= day <= 7: return "Week 1 (1-7)"
                elif 8 <= day <= 14: return "Week 2 (8-14)"
                elif 15 <= day <= 21: return "Week 3 (15-21)"
                elif 22 <= day <= 28: return "Week 4 (22-28)"
                elif day >= 29: return "Week 5 (29-31)"
                return np.nan
            final_df['Week'] = final_df.apply(assign_week_formatted, axis=1)
        else:
            final_df['Week'] = np.nan

        def compute_grouped_table(df_subset, group_col, selected_item="All", search_query=""):
            if group_col not in df_subset.columns:
                return pd.DataFrame()
            
            df_working = df_subset.dropna(subset=[group_col]).copy()
            df_working[group_col] = df_working[group_col].astype(str)
            
            if selected_item and selected_item != "All":
                df_working = df_working[df_working[group_col] == selected_item]

            if search_query and search_query.strip():
                df_working = df_working[df_working[group_col].str.contains(search_query.strip(), case=False, na=False)]
            
            if df_working.empty:
                return pd.DataFrame()

            grouped = df_working.groupby(group_col).agg(
                IMPRESSIONS=('_impressions', 'sum'),
                ATC=('_atc', 'sum'),
                ORDERS=('_orders', 'sum'),
                SPENDS=('_budget_consumed', 'sum'),
                SALES=('_sales', 'sum')
            ).reset_index()

            grouped['CPM'] = grouped.apply(lambda r: round((r['SPENDS'] / r['IMPRESSIONS']) * 1000) if r['IMPRESSIONS'] > 0 else 0, axis=1)
            grouped['ROAS'] = grouped.apply(lambda r: round(r['SALES'] / r['SPENDS'], 2) if r['SPENDS'] > 0 else 0.0, axis=1)
            grouped['ACOS'] = grouped.apply(lambda r: round((r['SPENDS'] / r['SALES']) * 100, 2) if r['SALES'] > 0 else 0.0, axis=1)

            display_name = group_col.upper()
            if group_col == 'Campaign Name': display_name = 'CAMPAIGN NAME'
            elif group_col == 'Ad Type Combined': display_name = 'MATCH / AD TYPE'

            grouped = grouped.rename(columns={group_col: display_name})
            col_order = [display_name, 'IMPRESSIONS', 'CPM', 'ATC', 'ORDERS', 'SPENDS', 'SALES', 'ROAS', 'ACOS']
            
            res_df = grouped.reindex(columns=col_order)
            res_df['SPENDS'] = res_df['SPENDS'].round(2)
            res_df['SALES'] = res_df['SALES'].round(2)
            return res_df

        MONTH_ORDER_DESC = {
            "JANUARY": 1, "FEBRUARY": 2, "MARCH": 3, "APRIL": 4,
            "MAY": 5, "JUNE": 6, "JULY": 7, "AUGUST": 8,
            "SEPTEMBER": 9, "OCTOBER": 10, "NOVEMBER": 11, "DECEMBER": 12
        }

        def get_calendar_months_desc(df_input):
            months = [str(x).strip().upper() for x in df_input['Month'].dropna().unique()]
            return sorted(months, key=lambda x: MONTH_ORDER_DESC.get(x, -1), reverse=True)

        def create_mom_comparison_table(df_input, entity_col):
            if entity_col not in df_input.columns:
                return pd.DataFrame()
            
            working_df = df_input.dropna(subset=[entity_col, 'Month']).copy()
            if working_df.empty:
                return pd.DataFrame()
            
            grouped = working_df.groupby([entity_col, 'Month']).agg(
                Impressions=('_impressions', 'sum'),
                ATC=('_atc', 'sum'),
                Orders=('_orders', 'sum'),
                Spends=('_budget_consumed', 'sum'),
                Sales=('_sales', 'sum')
            ).reset_index()

            grouped['CPM'] = grouped.apply(lambda r: round((r['Spends'] / r['Impressions']) * 1000) if r['Impressions'] > 0 else 0, axis=1)
            grouped['ROAS'] = grouped.apply(lambda r: round(r['Sales'] / r['Spends'], 2) if r['Spends'] > 0 else 0.0, axis=1)
            grouped['ACOS'] = grouped.apply(lambda r: round((r['Spends'] / r['Sales']) * 100, 2) if r['Sales'] > 0 else 0.0, axis=1)

            pivot_df = grouped.pivot(
                index=entity_col,
                columns='Month',
                values=['Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']
            )

            metrics_order = ['Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']
            all_months = get_calendar_months_desc(df_input)
            
            pivot_df = pivot_df.reorder_levels([1, 0], axis=1)
            sorted_cols = pd.MultiIndex.from_product([all_months, metrics_order], names=['Month', 'Metric'])
            pivot_df = pivot_df.reindex(columns=sorted_cols).fillna(0)

            grand_total_series = {}
            for month in all_months:
                month_df = working_df[working_df['Month'] == month]
                total_imp = month_df['_impressions'].sum()
                total_atc = month_df['_atc'].sum()
                total_orders = month_df['_orders'].sum()
                total_spends = round(month_df['_budget_consumed'].sum(), 2)
                total_sales = round(month_df['_sales'].sum(), 2)

                total_cpm = round((total_spends / total_imp) * 1000) if total_imp > 0 else 0
                total_roas = round(total_sales / total_spends, 2) if total_spends > 0 else 0.0
                total_acos = round((total_spends / total_sales) * 100, 2) if total_sales > 0 else 0.0

                grand_total_series[(month, 'Impressions')] = total_imp
                grand_total_series[(month, 'CPM')] = total_cpm
                grand_total_series[(month, 'ATC')] = total_atc
                grand_total_series[(month, 'Orders')] = total_orders
                grand_total_series[(month, 'Spends')] = total_spends
                grand_total_series[(month, 'Sales')] = total_sales
                grand_total_series[(month, 'ROAS')] = total_roas
                grand_total_series[(month, 'ACOS')] = total_acos

            pivot_df.loc['Grand Total'] = grand_total_series
            pivot_df.index.name = entity_col
            return pivot_df

        def create_monthly_summary_table(df_input):
            working_df = df_input.dropna(subset=['Month']).copy()
            if working_df.empty:
                return pd.DataFrame()

            monthly_agg = working_df.groupby('Month').agg(
                Impressions=('_impressions', 'sum'),
                ATC=('_atc', 'sum'),
                Orders=('_orders', 'sum'),
                Spends=('_budget_consumed', 'sum'),
                Sales=('_sales', 'sum')
            ).reset_index()

            monthly_agg['CPM'] = monthly_agg.apply(lambda r: round((r['Spends'] / r['Impressions']) * 1000) if r['Impressions'] > 0 else 0, axis=1)
            monthly_agg['ROAS'] = monthly_agg.apply(lambda r: round(r['Sales'] / r['Spends'], 2) if r['Spends'] > 0 else 0.0, axis=1)
            monthly_agg['ACOS'] = monthly_agg.apply(lambda r: round((r['Spends'] / r['Sales']) * 100, 2) if r['Sales'] > 0 else 0.0, axis=1)

            monthly_agg['Spends'] = monthly_agg['Spends'].round(2)
            monthly_agg['Sales'] = monthly_agg['Sales'].round(2)

            all_months = sorted(
                [str(x).strip().upper() for x in df_input['Month'].dropna().unique()],
                key=lambda x: MONTH_ORDER_DESC.get(x, 999)
            )
            month_order = {m: i for i, m in enumerate(all_months)}
            monthly_agg['month_order'] = monthly_agg['Month'].map(
                lambda x: month_order.get(str(x).strip().upper(), 999)
            )
            monthly_agg = monthly_agg.sort_values('month_order').drop(columns=['month_order'])

            col_order = ['Month', 'Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']
            monthly_agg = monthly_agg[col_order]

            if len(monthly_agg) >= 2:
                prev_row = monthly_agg.iloc[-2]
                curr_row = monthly_agg.iloc[-1]

                pct_row = {'Month': 'Percentage %'}
                for metric in ['Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']:
                    prev_val = prev_row[metric]
                    curr_val = curr_row[metric]
                    if prev_val > 0:
                        pct_change = round(((curr_val - prev_val) / prev_val) * 100)
                        pct_row[metric] = f"{pct_change}%" if pct_change <= 0 else f"+{pct_change}%"
                    else:
                        pct_row[metric] = "0%"

                monthly_agg = pd.concat([monthly_agg, pd.DataFrame([pct_row])], ignore_index=True)

            monthly_agg['ACOS'] = monthly_agg['ACOS'].apply(lambda v: f"{v:.2f}%" if isinstance(v, (int, float)) else str(v))
            return monthly_agg

        def _format_dashboard_value(value, column_name=""):
            if pd.isna(value):
                return value

            col = str(column_name).strip().upper()
            is_percentage = ('%' in col) or ('ACOS' in col)

            if isinstance(value, (int, np.integer)) and not isinstance(value, bool):
                return f"{int(value):,}" if not is_percentage else f"{int(value)}%"

            if isinstance(value, (float, np.floating)) and np.isfinite(value):
                if is_percentage:
                    return f"{value:.2f}%"
                if float(value).is_integer():
                    return f"{int(value):,}"
                return f"{value:,.2f}"

            return value

        def format_dashboard_dataframe(df):
            out = df.copy()
            if isinstance(out.columns, pd.MultiIndex):
                for col in out.columns:
                    out[col] = out[col].map(lambda v: _format_dashboard_value(v, f"{col[0]} {col[1]}"))
            else:
                for col in out.columns:
                    out[col] = out[col].map(lambda v: _format_dashboard_value(v, col))
            return out

        def render_unified_single_table(df_to_show, key_prefix="mom", expandable_col=None):
            if df_to_show is None or df_to_show.empty:
                return

            working_df = df_to_show.copy()

            bottom_rows = None
            if isinstance(working_df.index, pd.Index) and 'Grand Total' in working_df.index:
                main_df = working_df.drop('Grand Total')
                bottom_rows = working_df.loc[['Grand Total']]
            elif 'Month' in working_df.columns and 'Percentage %' in working_df['Month'].values:
                main_df = working_df[working_df['Month'] != 'Percentage %'].copy()
                bottom_rows = working_df[working_df['Month'] == 'Percentage %'].copy()
            else:
                main_df = working_df
                bottom_rows = None

            sort_cols = [str(c) for c in main_df.columns] if not isinstance(main_df.columns, pd.MultiIndex) else [f"{c[0]} - {c[1]}" for c in main_df.columns]
            c_sort1, c_sort2 = st.columns([3, 1])
            with c_sort1:
                selected_sort_col = st.selectbox("Sort Table Column:", ["Default Order"] + sort_cols, key=f"{key_prefix}_sort_col")
            with c_sort2:
                sort_order = st.radio("Order:", ["Ascending", "Descending"], key=f"{key_prefix}_sort_ord", horizontal=True)

            if selected_sort_col != "Default Order":
                asc = (sort_order == "Ascending")
                if isinstance(main_df.columns, pd.MultiIndex):
                    idx_match = sort_cols.index(selected_sort_col)
                    col_key = main_df.columns[idx_match]
                    main_df = main_df.sort_values(by=col_key, ascending=asc)
                else:
                    main_df = main_df.sort_values(by=selected_sort_col, ascending=asc)

            if bottom_rows is not None and not bottom_rows.empty:
                unified_df = pd.concat([main_df, bottom_rows])
            else:
                unified_df = main_df

            def highlight_percentage_cells(val):
                val_str = str(val).strip()
                if '%' in val_str:
                    if val_str.startswith('-'):
                        return 'background-color: #f8d7da; color: #721c24; font-weight: bold;'
                    elif val_str.startswith('+'):
                        return 'background-color: #d4edda; color: #155724; font-weight: bold;'
                return ''

            col_config = {}
            if expandable_col:
                col_config[expandable_col] = st.column_config.TextColumn(
                    expandable_col,
                    help="Click edge to expand full keyword string",
                    width="large"
                )

            if hasattr(unified_df.style, "map"):
                styled_unified_df = unified_df.style.map(highlight_percentage_cells)
            else:
                styled_unified_df = unified_df.style.applymap(highlight_percentage_cells)

            display_df = format_dashboard_dataframe(unified_df)
            st.dataframe(
                display_df.style.map(highlight_percentage_cells) if hasattr(display_df.style, "map") else display_df.style.applymap(highlight_percentage_cells),
                use_container_width=True,
                hide_index=False if isinstance(unified_df.index, pd.MultiIndex) or unified_df.index.name else True,
                column_config=col_config,
                key=f"{key_prefix}_single_unified_grid"
            )

        # Dashboard Filters
        st.markdown("### 🔍 Global Dashboard Filters")
        available_months = ["All Months"] + get_calendar_months_desc(final_df)
        selected_month = st.selectbox("Select Month Across Dashboard (Excluding Comparison Tables)", available_months)

        filtered_df = final_df.copy()
        if selected_month != "All Months":
            filtered_df = filtered_df[filtered_df['Month'] == selected_month]

        total_impressions = filtered_df['_impressions'].sum()
        total_sales = round(filtered_df['_sales'].sum(), 2)
        total_orders = filtered_df['_orders'].sum()
        total_atc = filtered_df['_atc'].sum()
        total_budget = round(filtered_df['_budget_consumed'].sum(), 2)
        overall_roas = round((total_sales / total_budget), 2) if total_budget > 0 else 0.0

        st.markdown("### 📈 Overall Campaign Performance Dashboard")
        
        row1_col1, row1_col2, row1_col3 = st.columns(3)
        with row1_col1: st.metric("Total Impressions", f"{int(total_impressions):,}")
        with row1_col2: st.metric("Total Sales", f"₹{total_sales:,.2f}")
        with row1_col3: st.metric("Total Budget Consumed", f"₹{total_budget:,.2f}")

        row2_col1, row2_col2, row2_col3 = st.columns(3)
        with row2_col1: st.metric("Overall RoAS", f"{overall_roas:.2f}x")
        with row2_col2: st.metric("Total Orders", f"{int(total_orders):,}")
        with row2_col3: st.metric("Total Add To Cart", f"{int(total_atc):,}")

        st.divider()

        st.markdown("### 📑 Navigation & Performance Breakdown")
        main_tab1, main_tab2, main_tab3, main_tab4, main_tab5, main_tab6, main_tab7, main_tab8 = st.tabs([
            "📄 Raw Files Preview",
            "📌 Consolidated Master",
            "📊 Comparison Tables (MoM)",
            "📈 Interactive Trend Analytics",
            "🎯 Campaign Performance", 
            "📢 Ad Type Performance",
            "🔎 Keyword Performance",
            "📅 Weekly Trend"
        ])

        with main_tab1:
            st.caption("Inspect individual sheets tab-by-tab for each uploaded file.")
            selected_file_name = st.selectbox("Select Uploaded File to Preview:", list(raw_files_dict.keys()))
            if selected_file_name:
                sheets = raw_files_dict[selected_file_name]
                selected_sheet = st.selectbox("Select Sheet Tab:", list(sheets.keys()))
                if selected_sheet:
                    st.dataframe(format_dashboard_dataframe(sheets[selected_sheet].head(100)), use_container_width=True, hide_index=True)

        with main_tab2:
            st.caption("Preview the combined dataset across all uploaded files.")
            preview_clean_df = final_df.drop(columns=['_impressions', '_direct_atc', '_indirect_atc', '_atc', '_direct_orders', '_indirect_orders', '_orders', '_direct_sales', '_indirect_sales', '_sales', '_budget_consumed', '_date_dt', 'Ad Type Combined'], errors='ignore')
            st.dataframe(format_dashboard_dataframe(preview_clean_df.head(100)), use_container_width=True, hide_index=True)

        with main_tab3:
            st.subheader("📊 Month-on-Month Comparison Tables")
            st.caption("View aggregated monthly summary & side-by-side MoM metrics for Campaign, Ad Type, Keywords, and Weeks.")

            monthly_summary_df = create_monthly_summary_table(final_df)
            camp_pivot = create_mom_comparison_table(final_df, 'Campaign Name') if 'Campaign Name' in final_df.columns else None
            ad_pivot = create_mom_comparison_table(final_df, 'Ad Type Combined') if 'Ad Type Combined' in final_df.columns else None
            
            kw_col = None
            for c in ['Search Term', 'Keyword', 'Targeting Value']:
                if c in final_df.columns:
                    kw_col = c
                    break
            kw_pivot = create_mom_comparison_table(final_df, kw_col) if kw_col else None
            week_pivot = create_mom_comparison_table(final_df, 'Week') if ('Week' in final_df.columns and final_df['Week'].notna().any()) else None

            mom_dict = {
                "Monthly_Summary": monthly_summary_df,
                "Campaign_MoM": camp_pivot,
                "AdType_MoM": ad_pivot,
                "Keyword_MoM": kw_pivot,
                "Weekly_MoM": week_pivot
            }
            all_pivots_bytes = convert_all_pivots_to_excel(mom_dict)

            st.download_button(
                label="📥 Download All MoM Comparison Tables (.xlsx)",
                data=all_pivots_bytes,
                file_name="All_MoM_Comparison_Tables_Combined.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="btn_dl_all_mom_pivots",
                type="primary"
            )

            st.divider()

            comp_sub_tab0, comp_sub_tab1, comp_sub_tab2, comp_sub_tab3, comp_sub_tab4 = st.tabs([
                "📅 Monthly Summary",
                "🎯 Campaign Comparison",
                "📢 Ad Type Comparison",
                "🔎 Keyword Comparison",
                "📅 Weekly Comparison"
            ])

            with comp_sub_tab0:
                st.markdown("#### Monthly Comparison Summary Table")
                if not monthly_summary_df.empty:
                    render_unified_single_table(monthly_summary_df, key_prefix="monthly_summary_tab")

                    excel_month_summary = style_and_export_pivot(monthly_summary_df, sheet_name="Monthly_Summary")
                    st.download_button(
                        label="📥 Download Monthly Summary Table (.xlsx)",
                        data=excel_month_summary,
                        file_name="Monthly_Summary_Comparison_Report.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="btn_dl_month_summary"
                    )

            with comp_sub_tab1:
                st.markdown("#### Campaign Month-on-Month Comparison Table")
                if camp_pivot is not None and not camp_pivot.empty:
                    render_unified_single_table(camp_pivot, key_prefix="camp_pivot_tab")
                    
                    excel_camp_pivot = style_and_export_pivot(camp_pivot, sheet_name="Campaign_MoM")
                    st.download_button(
                        label="📥 Download Campaign Comparison Table (.xlsx)",
                        data=excel_camp_pivot,
                        file_name="Campaign_MoM_Comparison_Report.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="btn_dl_camp_pivot"
                    )
                else:
                    st.info("No Campaign Name column found.")

            with comp_sub_tab2:
                st.markdown("#### Ad Type Month-on-Month Comparison Table")
                if ad_pivot is not None and not ad_pivot.empty:
                    render_unified_single_table(ad_pivot, key_prefix="ad_pivot_tab")
                    
                    excel_ad_pivot = style_and_export_pivot(ad_pivot, sheet_name="AdType_MoM")
                    st.download_button(
                        label="📥 Download Ad Type Comparison Table (.xlsx)",
                        data=excel_ad_pivot,
                        file_name="Ad_Type_MoM_Comparison_Report.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="btn_dl_ad_pivot"
                    )

            with comp_sub_tab3:
                st.markdown("#### Keyword / Search Term Month-on-Month Comparison Table")
                if kw_pivot is not None and not kw_pivot.empty:
                    render_unified_single_table(kw_pivot, key_prefix="kw_pivot_tab", expandable_col=kw_col)
                    
                    excel_kw_pivot = style_and_export_pivot(kw_pivot, sheet_name="Keyword_MoM")
                    st.download_button(
                        label="📥 Download Keyword Comparison Table (.xlsx)",
                        data=excel_kw_pivot,
                        file_name="Keyword_MoM_Comparison_Report.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="btn_dl_kw_pivot"
                    )
                else:
                    st.info("No Keyword or Search Term column found.")

            with comp_sub_tab4:
                st.markdown("#### Weekly Month-on-Month Comparison Table")
                if week_pivot is not None and not week_pivot.empty:
                    render_unified_single_table(week_pivot, key_prefix="week_pivot_tab")
                    
                    excel_week_pivot = style_and_export_pivot(week_pivot, sheet_name="Weekly_MoM")
                    st.download_button(
                        label="📥 Download Weekly Comparison Table (.xlsx)",
                        data=excel_week_pivot,
                        file_name="Weekly_MoM_Comparison_Report.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="btn_dl_week_pivot"
                    )

        with main_tab4:
            st.subheader("📈 Interactive Multi-Metric Trend Analytics")
            st.caption("Select multiple months and metrics to compare performance across time with smooth curved lines overlaying metric pillar columns.")

            all_df_months = get_calendar_months_desc(final_df)
            selected_trend_months = st.multiselect(
                "Select Months to Include in Trend Analysis:",
                options=all_df_months,
                default=all_df_months
            )

            metric_map = {
                'Sales (₹)': '_sales',
                'Spends (₹)': '_budget_consumed',
                'ROAS': 'ROAS',
                'Orders': '_orders',
                'Add To Cart (ATC)': '_atc',
                'Impressions': '_impressions',
                'ACOS (%)': 'ACOS',
                'CPM (₹)': 'CPM'
            }

            selected_trend_metrics = st.multiselect(
                "Select Metrics to Display on Trend Graph:",
                options=list(metric_map.keys()),
                default=['Sales (₹)', 'Spends (₹)', 'ROAS']
            )

            if selected_trend_months and selected_trend_metrics:
                trend_df = final_df[final_df['Month'].isin(selected_trend_months)].copy()

                monthly_summary = trend_df.groupby('Month').agg(
                    _impressions=('_impressions', 'sum'),
                    _atc=('_atc', 'sum'),
                    _orders=('_orders', 'sum'),
                    _budget_consumed=('_budget_consumed', 'sum'),
                    _sales=('_sales', 'sum')
                ).reset_index()

                monthly_summary['CPM'] = monthly_summary.apply(lambda r: round((r['_budget_consumed'] / r['_impressions']) * 1000) if r['_impressions'] > 0 else 0, axis=1)
                monthly_summary['ROAS'] = monthly_summary.apply(lambda r: round(r['_sales'] / r['_budget_consumed'], 2) if r['_budget_consumed'] > 0 else 0.0, axis=1)
                monthly_summary['ACOS'] = monthly_summary.apply(lambda r: round((r['_budget_consumed'] / r['_sales']) * 100, 2) if r['_sales'] > 0 else 0.0, axis=1)
                
                monthly_summary['_budget_consumed'] = monthly_summary['_budget_consumed'].round(2)
                monthly_summary['_sales'] = monthly_summary['_sales'].round(2)

                ordered_trend_months = get_calendar_months_desc(trend_df)
                month_order = {m: i for i, m in enumerate(ordered_trend_months)}
                monthly_summary['month_idx'] = monthly_summary['Month'].map(
                    lambda x: month_order.get(str(x).strip().upper(), 99)
                )
                monthly_summary = monthly_summary.sort_values('month_idx').drop(columns=['month_idx'])

                st.markdown("#### 📊 Selected Months Data Summary")
                disp_summary = monthly_summary.rename(columns={
                    'Month': 'MONTH',
                    '_impressions': 'IMPRESSIONS',
                    '_atc': 'ATC',
                    '_orders': 'ORDERS',
                    '_budget_consumed': 'SPENDS',
                    '_sales': 'SALES'
                })
                st.dataframe(format_dashboard_dataframe(disp_summary), use_container_width=True, hide_index=True)

                st.markdown("#### 📉 Curved Trend Line & Pillar Combination Graph")
                
                fig_trend = make_subplots(specs=[[{"secondary_y": True}]])

                metric_colors = {
                    'Sales (₹)': '#5B8DB8',
                    'Spends (₹)': '#8CB6D3',
                    'ROAS': '#6E8FB5',
                    'Orders': '#6AAE9B',
                    'Add To Cart (ATC)': '#A18DB8',
                    'Impressions': '#7FA7A3',
                    'ACOS (%)': '#C28FA0',
                    'CPM (₹)': '#B39A70'
                }

                bar_metrics = {'Sales (₹)', 'Spends (₹)'}

                for metric_label in selected_trend_metrics:
                    col_key = metric_map[metric_label]
                    use_sec_y = metric_label in ['ROAS', 'ACOS (%)', 'CPM (₹)']
                    color = metric_colors.get(metric_label, '#6AA6CF')
                    values = monthly_summary[col_key]

                    if metric_label in bar_metrics:
                        fig_trend.add_trace(
                            go.Bar(
                                x=monthly_summary['Month'],
                                y=values,
                                name=metric_label,
                                marker=dict(color=color, opacity=0.78, line=dict(color='#FFFFFF', width=1)),
                                text=values.map(lambda v: _format_dashboard_value(v, metric_label)),
                                textposition='outside',
                                hovertemplate=f"<b>{metric_label}</b>: %{{y:,.2f}}<extra></extra>"
                            ),
                            secondary_y=use_sec_y
                        )
                    else:
                        fig_trend.add_trace(
                            go.Scatter(
                                x=monthly_summary['Month'],
                                y=values,
                                name=metric_label,
                                mode='lines+markers',
                                line=dict(shape='spline', width=3.5, color=color),
                                marker=dict(size=8, color=color, symbol='circle', line=dict(color='#FFFFFF', width=1.5)),
                                hovertemplate=f"<b>{metric_label}</b>: %{{y:,.2f}}<extra></extra>"
                            ),
                            secondary_y=use_sec_y
                        )

                fig_trend.update_layout(
                    title="<b>Multi-Metric Trend Analysis</b>",
                    template='plotly_white',
                    height=580,
                    hovermode='x unified',
                    barmode='group',
                    bargap=0.28,
                    plot_bgcolor='#FBFDFF',
                    paper_bgcolor='#FFFFFF',
                    font=dict(family='Arial, sans-serif', color='#29465B'),
                    legend=dict(orientation='h', yanchor='bottom', y=1.02, xanchor='right', x=1),
                    xaxis=dict(title='Months', showgrid=False, categoryorder='array', categoryarray=monthly_summary['Month'].tolist()),
                    yaxis=dict(title='Sales / Spends / Volume', showgrid=True, gridcolor='#E8F0F6', zeroline=False),
                    yaxis2=dict(title='Ratios / Rates', overlaying='y', side='right', showgrid=False, zeroline=False)
                )

                st.plotly_chart(fig_trend, use_container_width=True)
            else:
                st.info("Please select at least one Month and one Metric to display the trend analysis.")

        with main_tab5:
            if 'Campaign Name' in filtered_df.columns:
                campaign_options = ["All"] + sorted([str(x) for x in filtered_df['Campaign Name'].dropna().unique()])
                selected_campaign = st.selectbox("Select Campaign:", campaign_options, key="camp_single_filt")
                campaign_df = compute_grouped_table(filtered_df, 'Campaign Name', selected_campaign)
                if not campaign_df.empty:
                    st.dataframe(format_dashboard_dataframe(campaign_df), use_container_width=True, hide_index=True)

        with main_tab6:
            if 'Ad Type Combined' in filtered_df.columns:
                adtype_options = ["All"] + sorted([str(x) for x in filtered_df['Ad Type Combined'].dropna().unique()])
                selected_adtype = st.selectbox("Select Ad Type:", adtype_options, key="ad_single_filt")
                adtype_df = compute_grouped_table(filtered_df, 'Ad Type Combined', selected_adtype)
                if not adtype_df.empty:
                    st.dataframe(format_dashboard_dataframe(adtype_df), use_container_width=True, hide_index=True)

        with main_tab7:
            kw_col = None
            for c in ['Search Term', 'Keyword', 'Targeting Value']:
                if c in filtered_df.columns:
                    kw_col = c
                    break
            if kw_col:
                st.markdown("#### 🔎 Keyword Performance & Search Filter")
                search_box_val = st.text_input("🔍 Search by Keyword (e.g. crack):", "", key="keyword_text_search")
                
                kw_options = ["All"] + sorted([str(x) for x in filtered_df[kw_col].dropna().unique()])
                selected_kw = st.selectbox(f"Select specific {kw_col}:", kw_options, key="kw_single_filt")
                
                effective_item = selected_kw if selected_kw != "All" else "All"
                search_df = compute_grouped_table(filtered_df, kw_col, effective_item, search_query=search_box_val)
                
                if not search_df.empty:
                    st.dataframe(
                        search_df,
                        use_container_width=True,
                        hide_index=True,
                        column_config={
                            kw_col.upper(): st.column_config.TextColumn(
                                kw_col.upper(),
                                width="large"
                            )
                        }
                    )
                else:
                    st.info("No matching keywords found for your search query.")

        with main_tab8:
            if 'Week' in filtered_df.columns and filtered_df['Week'].notna().any():
                weekly_df = compute_grouped_table(filtered_df, 'Week', "All")
                st.dataframe(format_dashboard_dataframe(weekly_df), use_container_width=True, hide_index=True)

        st.divider()

        st.subheader("💾 Download Consolidated Master Workbook")
        buffer_multi = io.BytesIO()
        with pd.ExcelWriter(buffer_multi, engine='openpyxl') as writer:
            for raw_tab_name, df_list in consolidated_raw_tabs.items():
                combined_raw_tab_df = pd.concat(df_list, ignore_index=True)
                clean_sheet_name = raw_tab_name[:31]
                combined_raw_tab_df.to_excel(writer, sheet_name=clean_sheet_name, index=False)
            
            master_export_df = final_df.drop(columns=['_impressions', '_direct_atc', '_indirect_atc', '_atc', '_direct_orders', '_indirect_orders', '_orders', '_direct_sales', '_indirect_sales', '_sales', '_budget_consumed', '_date_dt', 'Ad Type Combined'], errors='ignore')
            master_export_df.to_excel(writer, sheet_name='Consolidated_Master', index=False)

        buffer_multi.seek(0)
        st.download_button(
            label="📥 Download Complete Excel Workbook (Consolidated Raw Tabs + Final Master Tab)",
            data=buffer_multi,
            file_name="Blinkit_Consolidated_Master_Report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
